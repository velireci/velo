"""
Merge the owner lookup into the call list.

Takes leads_with_owners.csv (written by fetch_owners.py) and fills
owner_name into column A and owner_mailing into column D of the 'Call List'
sheet, matching rows on business_name. Everything else in the workbook -
formulas, conditional formatting, the drop-downs on the call-tracking
columns, the Dashboard and Rejected sheets - is left alone.

    python merge_owners.py                      # edits the workbook in place
    python merge_owners.py --out filled.xlsx    # writes a copy instead
    python merge_owners.py --dry-run            # report only, write nothing
"""

import argparse
import re
import sys
import unicodedata

import openpyxl
import pandas as pd

WORKBOOK = 'tampa_lodging_call_list.xlsx'
OWNERS = 'leads_with_owners.csv'
SHEET = 'Call List'
KEY_COL = 'H'            # Business Name
NAME_COL = 'A'           # Owner Name
MAIL_COL = 'D'           # Owner Mailing
HEADER_ROW = 1


def norm(s):
    """Fold case, curly quotes and runs of whitespace so keys compare equal."""
    if s is None:
        return ''
    s = unicodedata.normalize('NFKC', str(s))
    s = s.replace('’', "'").replace('‘', "'")
    return re.sub(r'\s+', ' ', s).strip().casefold()


def check_headers(ws):
    """Refuse to write into columns that are not the ones we think they are."""
    expected = {KEY_COL: 'business name',
                NAME_COL: 'owner name',
                MAIL_COL: 'owner mailing'}
    bad = {col: ws[f'{col}{HEADER_ROW}'].value
           for col, want in expected.items()
           if norm(ws[f'{col}{HEADER_ROW}'].value) != want}
    if bad:
        raise SystemExit(
            'column layout has changed - not writing. Expected '
            + ', '.join(f'{c}={w!r}' for c, w in expected.items())
            + '; found ' + ', '.join(f'{c}={v!r}' for c, v in bad.items()))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--workbook', default=WORKBOOK)
    ap.add_argument('--owners', default=OWNERS)
    ap.add_argument('--out', default=None, help='write here instead of in place')
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()

    owners = pd.read_csv(args.owners)
    for col in ('business_name', 'owner_name', 'owner_mailing'):
        if col not in owners.columns:
            raise SystemExit(f'{args.owners} has no {col} column')

    lookup = {}
    for r in owners.itertuples(index=False):
        key = norm(r.business_name)
        if key in lookup:
            print(f'!! duplicate business_name in {args.owners}: {r.business_name}')
        lookup[key] = (r.owner_name, r.owner_mailing)

    wb = openpyxl.load_workbook(args.workbook)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f'{args.workbook} has no {SHEET!r} sheet')
    ws = wb[SHEET]
    check_headers(ws)

    filled = blank = overwritten = 0
    unmatched = []
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        business = ws[f'{KEY_COL}{row}'].value
        if business is None or str(business).strip() == '':
            continue
        hit = lookup.get(norm(business))
        if hit is None:
            unmatched.append(business)
            continue
        name, mailing = hit
        name = '' if pd.isna(name) else str(name).strip()
        mailing = '' if pd.isna(mailing) else str(mailing).strip()
        if not name and not mailing:
            blank += 1
            continue
        for col, value in ((NAME_COL, name), (MAIL_COL, mailing)):
            cell = ws[f'{col}{row}']
            if cell.value not in (None, '') and str(cell.value).strip() != value:
                overwritten += 1
            cell.value = value or None
        filled += 1

    used = {norm(b) for b in
            (ws[f'{KEY_COL}{r}'].value for r in range(HEADER_ROW + 1, ws.max_row + 1))
            if b}
    missing_from_sheet = [r.business_name for r in owners.itertuples(index=False)
                          if norm(r.business_name) not in used]

    print(f'filled {filled} rows in {SHEET} '
          f'(columns {NAME_COL} and {MAIL_COL})')
    if overwritten:
        print(f'!! replaced {overwritten} cells that already held a different value')
    if blank:
        print(f'   {blank} leads had no owner in {args.owners} - left blank')
    if unmatched:
        print(f'!! {len(unmatched)} sheet rows had no owner record: '
              + ', '.join(map(str, unmatched[:10]))
              + (' ...' if len(unmatched) > 10 else ''))
    if missing_from_sheet:
        print(f'!! {len(missing_from_sheet)} owner records had no sheet row: '
              + ', '.join(map(str, missing_from_sheet[:10]))
              + (' ...' if len(missing_from_sheet) > 10 else ''))

    if args.dry_run:
        print('dry run - workbook not written')
        return 0
    dest = args.out or args.workbook
    wb.save(dest)
    print(f'wrote {dest}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
